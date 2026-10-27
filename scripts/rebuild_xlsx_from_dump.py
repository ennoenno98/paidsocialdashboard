"""Rebuild a values-only ``Social_Media.xlsx`` from a Microsoft 365 connector text dump.

The connector renders each worksheet as tab-separated rows (``## Sheet: <name> — R rows × C columns``,
``[N empty rows]`` markers, a trailing ``Formulas:`` section). This script reconstructs the cell grid
so the parser can run when the binary workbook is not reachable. Numeric-looking text (``2.344``,
``2,30%``, ``€ 160,00``, ``€ -``) is converted to numbers, as in the real cells.

Usage: python scripts/rebuild_xlsx_from_dump.py DUMP.txt OUT.xlsx [--only "Sheet A" ...] [--skip ...] [--tsv Creators=creators.tsv]
(multiple dumps: run once per dump into different files, or concatenate the dumps first)
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from src.ingest.labels import parse_number  # noqa: E402

SHEET_RE = re.compile(r"^## Sheet: (.+?)(?: \[hidden\])? — (\d+) rows × (\d+) columns")
EMPTY_RE = re.compile(r"^\[(\d+) empty rows?\]$")
NUMERIC_RE = re.compile(r"^[\s€$\d.,%\-–]+$")


def to_cell(text: str):  # noqa: ANN201
    t = text.strip()
    if t == "":
        return None
    if t.upper().startswith("#"):
        return t  # #DIV/0! etc. stay as error text
    if (
        NUMERIC_RE.match(t)
        and any(ch.isdigit() for ch in t)
        or ("€" in t and t.replace("€", "").strip() in ("-", "–"))
    ):
        v = parse_number(t)
        if v == v:  # not NaN
            return v
    return text.rstrip("\n")


def split_sheets(lines: list[str]) -> dict[str, list[list[object]]]:
    sheets: dict[str, list[list[object]]] = {}
    name: str | None = None
    rows: list[list[object]] = []
    in_formulas = False
    for ln in lines:
        m = SHEET_RE.match(ln)
        if m:
            if name:
                sheets[name] = rows
            name, rows, in_formulas = m.group(1).strip(), [], False
            continue
        if name is None or in_formulas:
            continue
        if ln.startswith("Formulas:"):
            in_formulas = True
            continue
        e = EMPTY_RE.match(ln.strip())
        if e:
            rows.extend([[] for _ in range(int(e.group(1)))])
            continue
        rows.append([to_cell(c) for c in ln.rstrip("\n").split("\t")])
    if name:
        sheets[name] = rows
    return sheets


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("dump")
    ap.add_argument("out")
    ap.add_argument("--only", nargs="*", default=None)
    ap.add_argument("--skip", nargs="*", default=[])
    ap.add_argument(
        "--tsv", nargs="*", default=[], help="extra sheets as NAME=path.tsv (tab-separated, no formulas)"
    )
    a = ap.parse_args()
    lines = Path(a.dump).read_text(encoding="utf-8").splitlines()
    sheets = split_sheets(lines)
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        if (a.only and name not in a.only) or name in a.skip:
            continue
        ws = wb.create_sheet(name[:31])
        for r in rows:
            ws.append(r if r else [None])
        print(f"{name}: {len(rows)} rows")
    for spec in a.tsv:
        name, _, path = spec.partition("=")
        ws = wb.create_sheet(name[:31])
        for ln in Path(path).read_text(encoding="utf-8").splitlines():
            ws.append([to_cell(c) for c in ln.split("\t")] or [None])
        print(f"{name}: tsv")
    wb.save(a.out)
    print("written", a.out)


if __name__ == "__main__":
    main()
