from __future__ import annotations

import math

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.ingest.xlsx import handle_from_url
from tests.conftest import XLSX


def test_block_coverage_all_titles_mapped(parsed):
    cov = parsed.coverage()
    assert not cov.empty
    assert (cov["status"] != "unmapped_title").all(), cov[cov["status"] == "unmapped_title"]
    assert (cov["status"] != "unparsed").all()
    weekly = cov[cov["sheet"] == "2026 SM Weekly Report"]
    assert len(weekly) == 11
    assert set(weekly["title"]) >= {
        "Vegavero Instagram Italy",
        "Boosting Italy",
        "Boosting TikTok",
        "Jasnum Instagram Italy",
    }


def test_empty_blocks_are_no_data_not_invented(parsed, weekly_wide):
    cov = parsed.coverage()
    nd = cov[cov["status"] == "no_data"]
    assert {"wowtamins Instagram Germany", "wowtamins TikTok Germany", "Jasnum Instagram Italy"} <= set(
        nd["title"]
    )
    w = weekly_wide[weekly_wide["account"].eq("wowtamins/instagram/DE")]
    assert w.empty or w.select_dtypes("number").drop(columns=["year", "period"]).isna().all().all()


def test_weekly_metrics_and_null_cells(weekly_wide):
    w = weekly_wide[
        (weekly_wide["account"] == "vegavero/instagram/IT")
        & (weekly_wide["kind"] == "organic")
        & (weekly_wide["year"] == 2026)
    ]
    assert len(w) == 43
    assert {
        "reach",
        "views",
        "follows",
        "unfollows",
        "net_follows",
        "followers",
        "n_posts",
        "n_igs",
        "er_sheet",
    } <= set(w.columns)
    cw7 = w[w["period"] == 7].iloc[0]
    assert math.isnan(cw7["er_sheet"])  # '#DIV/0!'
    cw12 = w[w["period"] == 12].iloc[0]
    assert math.isnan(cw12["reach"]) and math.isnan(cw12["er_sheet"])  # '/'
    assert w["net_follows"].notna().sum() == 43


def test_boosting_tiktok_gaps_and_duplicates_summed(parsed, weekly_wide):
    """Header 'CW 01, CW 02, CW 03, None, CW 04, None, CW 05': gaps ffilled, dupes summed (spend), text first."""
    cov = parsed.coverage()
    row = cov[(cov["sheet"] == "2026 SM Weekly Report") & (cov["title"] == "Boosting TikTok")].iloc[0]
    assert row["n_periods"] == 7
    assert "duplicate/filled period columns for [3, 4]" in row["warnings"]
    # read the raw cells back from the workbook: CW 03 = cols D+E, CW 04 = cols F+G of the spend row
    ws = load_workbook(XLSX, read_only=True, data_only=True)["2026 SM Weekly Report"]
    rows = list(ws.iter_rows(min_row=row["row_start"], max_row=row["row_end"], values_only=True))
    spend_row = next(r for r in rows if r[0] == "AD Spent")
    post_row = next(r for r in rows if r[0] == "Boosted post")
    long = parsed.weekly[
        (parsed.weekly["block_title"] == "Boosting TikTok") & (parsed.weekly["year"] == 2026)
    ]
    spend = long[long["metric"] == "boost_spend"].set_index("period")
    assert spend.loc[3, "n_cells"] == 2 and spend.loc[3, "value"] == pytest.approx(
        spend_row[3] + spend_row[4]
    )
    assert spend.loc[4, "n_cells"] == 2 and spend.loc[4, "value"] == pytest.approx(
        spend_row[5] + spend_row[6]
    )
    assert spend.loc[1, "n_cells"] == 1 and spend.loc[1, "value"] == pytest.approx(spend_row[1])
    text = long[long["metric"] == "boosted_post"].set_index("period")
    assert text.loc[3, "value_text"] == str(post_row[3]).strip() and text["value"].isna().all()
    w = weekly_wide[
        (weekly_wide["account"] == "vanatari/tiktok/INT")
        & (weekly_wide["kind"] == "boosting")
        & (weekly_wide["year"] == 2026)
    ]
    assert set(w["period"]) == {1, 2, 3, 4, 5, 6, 7}


def test_boosting_labels_mapped(weekly_wide):
    b = weekly_wide[weekly_wide["kind"] == "boosting"]
    assert {
        "boost_spend",
        "boosted_post",
        "goal",
        "audience",
        "reach",
        "new_followers",
        "link_clicks",
    } <= set(b.columns)
    de = b[(b["account"] == "vegavero/instagram/DE") & (b["year"] == 2026)]
    assert de["boost_spend"].notna().any()  # 'Total Spent' synonym


def test_monthly_typos_and_june_start(parsed, monthly_wide):
    m = monthly_wide[(monthly_wide["year"] == 2026) & (monthly_wide["period_type"] == "month")]
    it = m[m["account"] == "vegavero/instagram/IT"]
    assert set(it["period"]) == set(range(1, 11))  # data Jan-Oct; 'Janury ', 'July ' parsed
    cov = parsed.coverage()
    assert (
        cov[(cov["sheet"] == "2026 SM Monthly Report") & (cov["title"] == "Vegavero Instagram Italy")].iloc[
            0
        ]["n_periods"]
        == 12
    )
    assert it[it["period"] == 1]["reach"].notna().all()
    tt = m[m["account"] == "vegavero/tiktok/IT"]
    assert set(tt["period"]) == set(range(6, 11))  # header None for Jan-May; data Jun-Oct
    assert (
        cov[(cov["sheet"] == "2026 SM Monthly Report") & (cov["title"] == "Vegavero TikTok Italy")].iloc[0][
            "n_periods"
        ]
        == 7
    )
    assert {"boost_spend", "net_follows", "followers"} <= set(m.columns)
    de_sept = m[(m["account"] == "vegavero/instagram/DE") & (m["period"] == 9)].iloc[0]
    assert de_sept["boost_spend"] == 1200


def test_post_by_post(parsed):
    p = parsed.posts
    assert set(p["account"]) == {"vegavero/instagram/IT", "vegavero/instagram/DE", "vanatari/tiktok/INT"}
    ig = p[(p["account"] == "vegavero/instagram/IT") & (p["year"] == 2026)]
    assert ig["slot"].max() <= 4 and ig["cw"].max() <= 40
    assert ig["post_name"].notna().all() and ig["format"].notna().all()
    missing = ig[(ig["cw"] == 6) & (ig["slot"] == 1)].iloc[0]
    assert math.isnan(missing["reach"])
    tt = p[(p["account"] == "vanatari/tiktok/INT") & (p["year"] == 2026)]
    assert tt["video_views"].notna().all() and tt["play_time_min"].notna().all()
    assert set(tt["play_time_min"].round(1)) <= {72.0, 3720.0, 45.0, 180.0, 750.0}
    assert tt["shares"].notna().all() and tt["reposts"].isna().all()
    cov = parsed.coverage()
    r = cov[(cov["sheet"] == "2026 Post by Post") & (cov["title"] == "Instagram Italy")].iloc[0]
    assert "unlabelled columns" in r["warnings"]  # orphan ER cells ignored


def test_boosting_susy(parsed):
    s = parsed.boosting_susy
    de = s[(s["account"] == "vegavero/instagram/DE") & (s["month"] == 9)].iloc[0]
    assert de["boost_spend_susy"] == 1620
    b = parsed.susy_budget.set_index("account")
    assert b.loc["vegavero/instagram/DE", "boost_budget"] == 4000
    assert (
        math.isnan(b.loc["wowtamins/tiktok/DE", "boost_budget"])
        and b.loc["wowtamins/tiktok/DE", "budget_raw"] == "justus"
    )


def test_creators_no_pii(parsed):
    c = parsed.creators
    assert "address" not in {x.lower() for x in c.columns}
    assert not any("address" in str(col).lower() for col in c.columns)
    joined = c.to_string()
    assert "Roma" not in joined and "Berlin" not in joined and "Musterstr" not in joined
    assert set(c.columns) >= {
        "name",
        "country",
        "followers",
        "field",
        "account_url",
        "handle",
        "platform",
        "deal",
        "code",
    }
    g = c[c["name"] == "Giulia Fit"].iloc[0]
    assert g["handle"] == "fit_giulia" and g["platform"] == "instagram" and g["code"] == "GIULIA10"
    lena = c[c["name"] == "Lena Yoga"].iloc[0]
    assert lena["handle"] == "yogamitlena" and lena["platform"] == "tiktok"
    cov = parsed.coverage()
    assert "PII columns never read: ['Address']" in cov[cov["sheet"] == "Creators"].iloc[0]["warnings"]


def test_handle_from_url():
    assert handle_from_url("https://www.instagram.com/Fit_Giulia/") == ("fit_giulia", "instagram")
    assert handle_from_url("https://www.tiktok.com/@yogamitlena?lang=de") == ("yogamitlena", "tiktok")
    assert handle_from_url("@marco") == ("marco", None)
    assert handle_from_url(None) == (None, None)


def test_2025_sheets_parsed(parsed):
    cov = parsed.coverage()
    assert (cov["sheet"] == "2025 SM Report").sum() == 11
    assert (cov["sheet"] == "2025 Post by post report").sum() == 3
    assert (parsed.weekly["year"] == 2025).any()
    assert isinstance(parsed.posts, pd.DataFrame) and (parsed.posts["year"] == 2025).any()
