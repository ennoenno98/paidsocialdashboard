"""Parser for the human-formatted ``Social_Media.xlsx`` workbook.

Every relevant sheet is a stack of *blocks*: a title row (col A), a period header row
(``CW 01 …`` or month names in cols B+), then label rows. This module turns each block into
tidy long-format frames and produces a block-coverage report so the result can be verified
against the sheet. Nothing outside the configured sheets is read; the Creators address
column is never selected.
"""

from __future__ import annotations

import math
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Literal

import pandas as pd
import yaml
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from src.ingest.labels import (
    TEXT_METRICS,
    canonical_metric,
    is_blank_row,
    norm_label,
    parse_cw,
    parse_month,
    parse_number,
    parse_play_time,
)
from src.models import AccountKey, BlockReport, Issue

PeriodType = Literal["cw", "month"]
ACCOUNT_COLS = ["brand", "platform", "market", "kind", "account"]
CONFIG_PATH = Path(__file__).resolve().parents[2] / "config" / "accounts.yaml"

_FORBIDDEN_CREATOR_HEADERS = re.compile(
    r"addr|adres|anschrift|street|stra(ß|ss)e|plz|zip|postal|city|stadt|\bort\b|e-?mail|phone|tel|iban",
    re.I,
)
_CREATOR_HEADERS: dict[str, str] = {
    "country": "country",
    "land": "country",
    "followers": "followers",
    "follower": "followers",
    "field": "field",
    "account": "account_url",
    "account (url)": "account_url",
    "url": "account_url",
    "deal": "deal",
    "name": "name",
    "code": "code",
    "discount code": "code",
    "platform": "platform_raw",
}


# --------------------------------------------------------------------------------------
# config
# --------------------------------------------------------------------------------------
def load_config(path: Path | str = CONFIG_PATH) -> dict[str, Any]:
    with open(path, encoding="utf-8") as fh:
        return yaml.safe_load(fh)


def block_accounts(config: dict[str, Any]) -> dict[str, AccountKey]:
    return {norm_label(k): AccountKey(**v) for k, v in config["blocks"].items()}


def resolve_account(title: Any, accounts: dict[str, AccountKey]) -> AccountKey | None:
    n = norm_label(title)
    if n in accounts:
        return accounts[n]
    n2 = re.sub(r"\s*\(.*?\)\s*", " ", n).strip()  # drop parenthesised remarks
    return accounts.get(n2)


# --------------------------------------------------------------------------------------
# sheet scanning primitives
# --------------------------------------------------------------------------------------
@dataclass
class _Period:
    col: int  # 1-based column index
    value: int
    ffilled: bool


@dataclass
class _Block:
    sheet: str
    title: str
    title_row: int  # 1-based excel row
    header_row: int
    period_type: PeriodType
    periods: list[_Period]
    body: list[tuple[int, list[Any]]] = field(default_factory=list)  # (excel_row, values)

    @property
    def row_end(self) -> int:
        return self.body[-1][0] if self.body else self.header_row


def sheet_rows(ws) -> list[list[Any]]:  # noqa: ANN001
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    while rows and is_blank_row(rows[-1]):
        rows.pop()
    return rows


def _parse_header(row: list[Any], ptype: PeriodType) -> list[_Period]:
    """Parse cols B+ of a header row; forward-fill gaps *after* the first label."""
    parse = parse_cw if ptype == "cw" else parse_month
    out: list[_Period] = []
    current: int | None = None
    for i, cell in enumerate(row[1:], start=2):
        if cell is None or (isinstance(cell, str) and not cell.strip()):
            if current is not None:
                out.append(_Period(i, current, True))
            continue
        if ptype == "cw" and isinstance(cell, (int, float)):
            current = None  # bare numbers are never CW labels (would be confused with data)
            continue
        v = parse(cell)
        if v is None:
            current = None  # e.g. "Total", "Budget", "Notes": stop filling
            continue
        current = v
        out.append(_Period(i, v, False))
    # only fill *gaps* between labels, never trailing columns after the last label
    while out and out[-1].ffilled:
        out.pop()
    return out


def _header_type(row: list[Any]) -> PeriodType | None:
    """A header row has >=2 labelled periods of one type (strings/dates only, not bare ints)."""
    if len(row) < 2:
        return None
    cells = [c for c in row[1:] if c is not None and not (isinstance(c, str) and not c.strip())]
    if not cells:
        return None
    n_cw = sum(1 for c in cells if isinstance(c, str) and parse_cw(c) is not None)
    n_m = sum(1 for c in cells if isinstance(c, (str, date, datetime)) and parse_month(c) is not None)
    if n_cw >= 2 and n_cw >= n_m:
        return "cw"
    if n_m >= 2:
        return "month"
    return None


def _is_title_row(row: list[Any]) -> bool:
    a = row[0] if row else None
    if not isinstance(a, str) or not a.strip():
        return False
    return all(c is None or isinstance(c, str) for c in row[1:])


def find_blocks(rows: list[list[Any]], sheet: str) -> list[_Block]:
    """Generic block scanner used by all report-style sheets."""
    n = len(rows)
    title_rows: set[int] = set()
    for i in range(n - 1):
        if _is_title_row(rows[i]) and _header_type(rows[i + 1]) is not None:
            title_rows.add(i)
    blocks: list[_Block] = []
    for i in sorted(title_rows):
        ptype = _header_type(rows[i + 1])
        assert ptype is not None
        blk = _Block(sheet, str(rows[i][0]).strip(), i + 1, i + 2, ptype, _parse_header(rows[i + 1], ptype))
        j = i + 2
        while j < n and j not in title_rows and not is_blank_row(rows[j]):
            blk.body.append((j + 1, rows[j]))
            j += 1
        blocks.append(blk)
    return blocks


def _cell_ref(col: int, row: int) -> str:
    return f"{get_column_letter(col)}{row}"


# --------------------------------------------------------------------------------------
# Weekly / Monthly report sheets
# --------------------------------------------------------------------------------------
def parse_report_sheet(
    rows: list[list[Any]],
    sheet: str,
    year: int,
    accounts: dict[str, AccountKey],
    expected: PeriodType | None = None,
) -> tuple[pd.DataFrame, list[BlockReport], list[Issue]]:
    """Weekly/monthly KPI blocks -> long frame (one row per account/period/metric)."""
    records: list[dict[str, Any]] = []
    reports: list[BlockReport] = []
    issues: list[Issue] = []

    for blk in find_blocks(rows, sheet):
        acc = resolve_account(blk.title, accounts)
        rep = BlockReport(
            sheet=sheet,
            title=blk.title,
            row_start=blk.title_row,
            row_end=blk.row_end,
            period_type=blk.period_type,
            year=year,
            account=acc.key if acc else None,
            kind=acc.kind if acc else None,
        )
        if expected and blk.period_type != expected:
            rep.warnings.append(f"header is {blk.period_type}, sheet expected {expected}")
        if acc is None:
            rep.status = "unmapped_title"
            issues.append(
                Issue(
                    source=sheet,
                    severity="error",
                    where=f"row {blk.title_row}",
                    message=f"block title '{blk.title}' not in config/accounts.yaml",
                )
            )
            reports.append(rep)
            continue

        # which forward-filled columns actually carry data?
        used_cols = {p.col for p in blk.periods if not p.ffilled}
        for p in blk.periods:
            if p.ffilled and any(len(r) >= p.col and r[p.col - 1] is not None for _, r in blk.body):
                used_cols.add(p.col)
        periods = [p for p in blk.periods if p.col in used_cols]
        labelled_cols = {p.col for p in blk.periods}
        n_values = 0
        seen_metrics: set[str] = set()
        for excel_row, r in blk.body:
            label = r[0] if r else None
            if label is None or (isinstance(label, str) and not label.strip()):
                continue
            metric = canonical_metric(label)
            if metric is None:
                metric = "x_" + re.sub(r"\W+", "_", norm_label(label)).strip("_")
                rep.warnings.append(f"unknown metric label '{label}' (row {excel_row}) kept as {metric}")
            seen_metrics.add(metric)
            # data sitting in columns without a period label
            stray = [
                i + 1 for i, c in enumerate(r) if i >= 1 and c is not None and (i + 1) not in labelled_cols
            ]
            if stray:
                rep.warnings.append(
                    f"row {excel_row}: values in unlabelled columns {[get_column_letter(c) for c in stray]} ignored"
                )
            for p in periods:
                raw = r[p.col - 1] if len(r) >= p.col else None
                if raw is None or (isinstance(raw, str) and not raw.strip()):
                    continue
                is_text = metric in TEXT_METRICS
                num = math.nan if is_text else parse_number(raw)
                txt = str(raw).strip() if (is_text or (isinstance(raw, str) and math.isnan(num))) else None
                if not is_text and not math.isnan(num):
                    n_values += 1
                records.append(
                    {
                        "sheet": sheet,
                        "block_title": blk.title,
                        "brand": acc.brand,
                        "platform": acc.platform,
                        "market": acc.market,
                        "kind": acc.kind,
                        "account": acc.key,
                        "year": year,
                        "period_type": blk.period_type,
                        "period": p.value,
                        "metric": metric,
                        "value": num,
                        "value_text": txt,
                        "cell": _cell_ref(p.col, excel_row),
                        "ffilled": p.ffilled,
                    }
                )
        rep.n_periods = len({p.value for p in periods})
        rep.n_metrics = len(seen_metrics)
        rep.n_values = n_values
        if n_values == 0:
            rep.status = "no_data"
        dup = {p.value for p in periods if sum(1 for q in periods if q.value == p.value) > 1}
        if dup:
            how = "summed" if acc.kind == "boosting" else "first kept"
            rep.warnings.append(
                f"duplicate/filled period columns for {sorted(dup)} -> numeric {how}, text first"
            )
        reports.append(rep)

    long = pd.DataFrame.from_records(
        records,
        columns=[
            "sheet",
            "block_title",
            "brand",
            "platform",
            "market",
            "kind",
            "account",
            "year",
            "period_type",
            "period",
            "metric",
            "value",
            "value_text",
            "cell",
            "ffilled",
        ],
    )
    if long.empty:
        return long, reports, issues
    long = _aggregate_duplicates(long)
    return long, reports, issues


def _aggregate_duplicates(long: pd.DataFrame) -> pd.DataFrame:
    keys = [
        "sheet",
        "block_title",
        "brand",
        "platform",
        "market",
        "kind",
        "account",
        "year",
        "period_type",
        "period",
        "metric",
    ]
    long = long.copy()
    long["is_boost"] = long["kind"].eq("boosting")

    def agg(g: pd.DataFrame) -> pd.Series:
        vals = g["value"].dropna()
        if g["is_boost"].iat[0]:
            value = vals.sum() if len(vals) else math.nan
        else:
            value = vals.iat[0] if len(vals) else math.nan
        txts = g["value_text"].dropna()
        return pd.Series(
            {
                "value": value,
                "value_text": txts.iat[0] if len(txts) else None,
                "cell": ",".join(g["cell"]),
                "n_cells": len(g),
            }
        )

    out = long.groupby(keys, sort=False, dropna=False).apply(agg, include_groups=False).reset_index()
    return out


def pivot_wide(long: pd.DataFrame) -> pd.DataFrame:
    """Long -> one row per account/period with metric columns (numeric + text)."""
    if long.empty:
        return pd.DataFrame(columns=ACCOUNT_COLS + ["year", "period_type", "period"])
    idx = ACCOUNT_COLS + ["year", "period_type", "period"]
    num = long[~long["metric"].isin(TEXT_METRICS)].pivot_table(
        index=idx, columns="metric", values="value", aggfunc="first"
    )
    txt = long[long["metric"].isin(TEXT_METRICS)]
    if not txt.empty:
        t = txt.pivot_table(index=idx, columns="metric", values="value_text", aggfunc="first")
        num = num.join(t, how="outer")
    wide = num.reset_index()
    wide.columns.name = None
    return wide.sort_values(idx).reset_index(drop=True)


# --------------------------------------------------------------------------------------
# Post by Post
# --------------------------------------------------------------------------------------
def _cell_getter(metric_rows: dict[str, tuple[int, list[Any]]], col: int):  # noqa: ANN202
    def cell(m: str) -> Any:
        rr = metric_rows.get(m)
        if rr is None or len(rr[1]) < col:
            return None
        return rr[1][col - 1]

    return cell


POST_COLS = [
    "sheet",
    "brand",
    "platform",
    "market",
    "kind",
    "account",
    "year",
    "cw",
    "slot",
    "post_name",
    "format",
    "reach",
    "views",
    "video_views",
    "play_time_min",
    "likes",
    "saves",
    "follows",
    "reposts",
    "shares",
    "comments",
    "er_sheet",
    "src_rows",
]


def parse_post_by_post(
    rows: list[list[Any]], sheet: str, year: int, accounts: dict[str, AccountKey]
) -> tuple[pd.DataFrame, list[BlockReport], list[Issue]]:
    """Blocks of repeating per-post row groups -> one row per (account, cw, slot)."""
    records: list[dict[str, Any]] = []
    reports: list[BlockReport] = []
    issues: list[Issue] = []
    for blk in find_blocks(rows, sheet):
        acc = resolve_account(blk.title, accounts)
        rep = BlockReport(
            sheet=sheet,
            title=blk.title,
            row_start=blk.title_row,
            row_end=blk.row_end,
            period_type=blk.period_type,
            year=year,
            account=acc.key if acc else None,
            kind=acc.kind if acc else None,
        )
        if blk.period_type != "cw":
            rep.status = "unparsed"
            rep.warnings.append("expected CW header")
            reports.append(rep)
            continue
        if acc is None:
            rep.status = "unmapped_title"
            issues.append(
                Issue(
                    source=sheet,
                    severity="error",
                    where=f"row {blk.title_row}",
                    message=f"block title '{blk.title}' not in config/accounts.yaml",
                )
            )
            reports.append(rep)
            continue
        periods = [p for p in blk.periods if not p.ffilled]
        labelled_cols = {p.col for p in periods}
        # split body into groups starting at each 'Post' label row
        groups: list[list[tuple[int, list[Any]]]] = []
        for excel_row, r in blk.body:
            if canonical_metric(r[0] if r else None) == "post_name":
                groups.append([])
            if not groups:
                rep.warnings.append(f"row {excel_row} before first 'Post' label ignored")
                continue
            groups[-1].append((excel_row, r))
        stray_cols: set[str] = set()
        n_values = 0
        for slot, grp in enumerate(groups, start=1):
            metric_rows: dict[str, tuple[int, list[Any]]] = {}
            for excel_row, r in grp:
                m = canonical_metric(r[0])
                if m is None:
                    rep.warnings.append(f"unknown label '{r[0]}' (row {excel_row}) in post group {slot}")
                    continue
                metric_rows.setdefault(m, (excel_row, r))
                for i, c in enumerate(r):
                    if i >= 1 and c is not None and (i + 1) not in labelled_cols:
                        stray_cols.add(get_column_letter(i + 1))
            for p in periods:
                cell = _cell_getter(metric_rows, p.col)
                name = cell("post_name")
                nums = {
                    m: parse_number(cell(m))
                    for m in (
                        "reach",
                        "views",
                        "video_views",
                        "likes",
                        "saves",
                        "follows",
                        "reposts",
                        "shares",
                        "comments",
                        "new_followers",
                        "er_sheet",
                    )
                }
                if nums["follows"] != nums["follows"] and not math.isnan(nums["new_followers"]):
                    nums["follows"] = nums["new_followers"]
                play = parse_play_time(cell("play_time"))
                has_data = any(
                    not math.isnan(v) for k, v in nums.items() if k != "er_sheet"
                ) or not math.isnan(play)
                if (name is None or not str(name).strip()) and not has_data:
                    continue
                n_values += sum(1 for v in nums.values() if not math.isnan(v))
                fmt = cell("format")
                records.append(
                    {
                        "sheet": sheet,
                        "brand": acc.brand,
                        "platform": acc.platform,
                        "market": acc.market,
                        "kind": acc.kind,
                        "account": acc.key,
                        "year": year,
                        "cw": p.value,
                        "slot": slot,
                        "post_name": str(name).strip() if name is not None else None,
                        "format": str(fmt).strip() if fmt is not None else None,
                        "reach": nums["reach"],
                        "views": nums["views"],
                        "video_views": nums["video_views"],
                        "play_time_min": play,
                        "likes": nums["likes"],
                        "saves": nums["saves"],
                        "follows": nums["follows"],
                        "reposts": nums["reposts"],
                        "shares": nums["shares"],
                        "comments": nums["comments"],
                        "er_sheet": nums["er_sheet"],
                        "src_rows": f"{grp[0][0]}-{grp[-1][0]}:{get_column_letter(p.col)}",
                    }
                )
        if stray_cols:
            rep.warnings.append(f"values in unlabelled columns {sorted(stray_cols)} ignored (orphan cells)")
        rep.n_periods = len({p.value for p in periods})
        rep.n_metrics = len(groups)
        rep.n_values = n_values
        if n_values == 0:
            rep.status = "no_data"
        reports.append(rep)
    posts = pd.DataFrame.from_records(records, columns=POST_COLS)
    return posts, reports, issues


# --------------------------------------------------------------------------------------
# Boosting Susy
# --------------------------------------------------------------------------------------
def parse_boosting_susy(
    rows: list[list[Any]], sheet: str, year: int, accounts: dict[str, AccountKey], budget_col: str = "E"
) -> tuple[pd.DataFrame, pd.DataFrame, list[BlockReport], list[Issue]]:
    """Manual spend table: title row, month row, spend row (+ budget in ``budget_col``)."""
    bcol = column_index_from_string(budget_col)
    spend_recs: list[dict[str, Any]] = []
    budget_recs: list[dict[str, Any]] = []
    reports: list[BlockReport] = []
    issues: list[Issue] = []
    n = len(rows)
    i = 0
    while i < n - 1:
        row = rows[i]
        if _is_title_row(row) and _header_type(rows[i + 1]) == "month":
            title = str(row[0]).strip()
            acc = resolve_account(title, accounts)
            hdr = rows[i + 1]
            periods = [p for p in _parse_header(hdr, "month") if not p.ffilled]
            spend_rows = []
            j = i + 2
            while (
                j < n
                and not is_blank_row(rows[j])
                and not (_is_title_row(rows[j]) and j + 1 < n and _header_type(rows[j + 1]) == "month")
            ):
                spend_rows.append((j + 1, rows[j]))
                j += 1
            rep = BlockReport(
                sheet=sheet,
                title=title,
                row_start=i + 1,
                row_end=j,
                period_type="month",
                year=year,
                account=acc.key if acc else None,
                kind="boosting",
            )
            if acc is None:
                rep.status = "unmapped_title"
                issues.append(
                    Issue(
                        source=sheet,
                        severity="error",
                        where=f"row {i + 1}",
                        message=f"account label '{title}' not in config/accounts.yaml",
                    )
                )
            else:
                n_values = 0
                for excel_row, r in spend_rows:
                    for p in periods:
                        raw = r[p.col - 1] if len(r) >= p.col else None
                        if raw is None:
                            continue
                        v = parse_number(raw)
                        if not math.isnan(v):
                            n_values += 1
                        elif isinstance(raw, str) and raw.strip():
                            rep.warnings.append(
                                f"{_cell_ref(p.col, excel_row)} is text '{raw.strip()}' -> NaN"
                            )
                        spend_recs.append(
                            {
                                "brand": acc.brand,
                                "platform": acc.platform,
                                "market": acc.market,
                                "kind": "boosting",
                                "account": acc.key,
                                "year": year,
                                "month": p.value,
                                "boost_spend_susy": v,
                                "cell": _cell_ref(p.col, excel_row),
                            }
                        )
                    # budget column
                    raw_b = r[bcol - 1] if len(r) >= bcol else None
                    hdr_b = hdr[bcol - 1] if len(hdr) >= bcol else None
                    if raw_b is not None and str(raw_b).strip():
                        if parse_month(hdr_b) is not None and isinstance(hdr_b, (str, date, datetime)):
                            rep.warnings.append(
                                f"budget col {budget_col} carries a month header ('{hdr_b}'); value read as spend, "
                                "not budget -> verify layout"
                            )
                        else:
                            bv = parse_number(raw_b)
                            if math.isnan(bv):
                                rep.warnings.append(
                                    f"budget cell {_cell_ref(bcol, excel_row)}='{raw_b}' is not a number -> NaN"
                                )
                            budget_recs.append(
                                {
                                    "brand": acc.brand,
                                    "platform": acc.platform,
                                    "market": acc.market,
                                    "kind": "boosting",
                                    "account": acc.key,
                                    "year": year,
                                    "boost_budget": bv,
                                    "budget_raw": str(raw_b).strip(),
                                    "cell": _cell_ref(bcol, excel_row),
                                }
                            )
                rep.n_periods = len(periods)
                rep.n_metrics = len(spend_rows)
                rep.n_values = n_values
                if n_values == 0:
                    rep.status = "no_data"
            reports.append(rep)
            i = j
        else:
            i += 1
    spend = pd.DataFrame.from_records(
        spend_recs, columns=ACCOUNT_COLS + ["year", "month", "boost_spend_susy", "cell"]
    )
    if not spend.empty:
        spend = spend.groupby(ACCOUNT_COLS + ["year", "month"], as_index=False).agg(
            boost_spend_susy=("boost_spend_susy", "sum"), cell=("cell", ",".join), n_cells=("cell", "size")
        )
        spend.loc[spend["n_cells"].eq(0), "boost_spend_susy"] = math.nan
    budget = pd.DataFrame.from_records(
        budget_recs, columns=ACCOUNT_COLS + ["year", "boost_budget", "budget_raw", "cell"]
    )
    return spend, budget, reports, issues


# --------------------------------------------------------------------------------------
# Creators (PII: address column is never selected)
# --------------------------------------------------------------------------------------
_HANDLE_RE = {
    "instagram": re.compile(r"instagram\.com/([A-Za-z0-9_.]+)", re.I),
    "tiktok": re.compile(r"tiktok\.com/@?([A-Za-z0-9_.]+)", re.I),
    "youtube": re.compile(r"youtube\.com/(?:@|c/|channel/|user/)?([A-Za-z0-9_.\-]+)", re.I),
}


def handle_from_url(url: Any) -> tuple[str | None, str | None]:
    """-> (handle, platform_raw). Bare '@handle' -> (handle, None)."""
    if url is None:
        return None, None
    s = str(url).strip()
    if not s:
        return None, None
    for plat, rx in _HANDLE_RE.items():
        m = rx.search(s)
        if m:
            return m.group(1).lower().rstrip("."), plat
    if s.startswith("@"):
        return s[1:].lower(), None
    if re.fullmatch(r"[A-Za-z0-9_.]+", s):
        return s.lower(), None
    return None, None


def parse_creators(ws) -> tuple[pd.DataFrame, list[BlockReport], list[Issue]]:  # noqa: ANN001
    issues: list[Issue] = []
    rows = sheet_rows(ws)
    header_idx: int | None = None
    colmap: dict[int, str] = {}
    for i, r in enumerate(rows[:30]):
        found = {}
        for ci, c in enumerate(r):
            if isinstance(c, str):
                n = norm_label(c)
                if n in _CREATOR_HEADERS and not _FORBIDDEN_CREATOR_HEADERS.search(n):
                    found[ci] = _CREATOR_HEADERS[n]
        if len(set(found.values())) >= 3:
            header_idx, colmap = i, found
            break
    rep = BlockReport(sheet=ws.title, title="Creators", row_start=(header_idx or 0) + 1, row_end=len(rows))
    if header_idx is None:
        rep.status = "unparsed"
        rep.warnings.append("header row with Country/Followers/Account/Code not found")
        issues.append(
            Issue(source=ws.title, severity="error", where="sheet", message="creators header not found")
        )
        return (
            pd.DataFrame(
                columns=[
                    "name",
                    "country",
                    "followers",
                    "field",
                    "account_url",
                    "handle",
                    "platform",
                    "platform_raw",
                    "deal",
                    "code",
                ]
            ),
            [rep],
            issues,
        )
    dropped = [
        str(c) for c in rows[header_idx] if isinstance(c, str) and _FORBIDDEN_CREATOR_HEADERS.search(c)
    ]
    if dropped:
        rep.warnings.append(f"PII columns never read: {dropped}")
    recs: list[dict[str, Any]] = []
    for r in rows[header_idx + 1 :]:
        if is_blank_row(r):
            continue
        rec: dict[str, Any] = {v: None for v in _CREATOR_HEADERS.values()}
        for ci, name in colmap.items():
            if ci < len(r):
                rec[name] = r[ci]
        if all(rec.get(k) in (None, "") for k in ("name", "account_url", "code")):
            continue
        rec["followers"] = parse_number(rec.get("followers"))
        rec["handle"], plat = handle_from_url(rec.get("account_url"))
        rec["platform_raw"] = rec.get("platform_raw") or plat
        rec["platform"] = (
            plat
            if plat in ("instagram", "tiktok")
            else (
                str(rec["platform_raw"]).strip().lower()
                if str(rec.get("platform_raw") or "").strip().lower() in ("instagram", "tiktok")
                else None
            )
        )
        rec["code"] = str(rec["code"]).strip().upper() if rec.get("code") not in (None, "") else None
        for k in ("name", "country", "field", "deal", "account_url"):
            rec[k] = str(rec[k]).strip() if rec.get(k) not in (None, "") else None
        recs.append(rec)
    df = pd.DataFrame.from_records(
        recs,
        columns=[
            "name",
            "country",
            "followers",
            "field",
            "account_url",
            "handle",
            "platform",
            "platform_raw",
            "deal",
            "code",
        ],
    )
    rep.n_values = len(df)
    rep.n_metrics = len(colmap)
    if df.empty:
        rep.status = "no_data"
    if not df.empty and df["handle"].isna().any():
        issues.append(
            Issue(
                source=ws.title,
                severity="warn",
                where="Account column",
                message=f"{int(df['handle'].isna().sum())} creators without a parsable handle/URL",
            )
        )
    return df, [rep], issues


# --------------------------------------------------------------------------------------
# workbook entry point
# --------------------------------------------------------------------------------------
@dataclass
class ParseResult:
    weekly: pd.DataFrame
    monthly: pd.DataFrame
    posts: pd.DataFrame
    boosting_susy: pd.DataFrame
    susy_budget: pd.DataFrame
    creators: pd.DataFrame
    blocks: list[BlockReport]
    issues: list[Issue]
    source: str

    def coverage(self) -> pd.DataFrame:
        df = pd.DataFrame([b.model_dump() for b in self.blocks])
        if df.empty:
            return df
        df["warnings"] = df["warnings"].map(lambda w: " | ".join(w))
        df["year"] = df["year"].astype("Int64")
        return df

    def issues_df(self) -> pd.DataFrame:
        return pd.DataFrame(
            [i.model_dump() for i in self.issues], columns=["source", "severity", "where", "message"]
        )

    def frames(self) -> dict[str, pd.DataFrame]:
        return {
            "weekly": self.weekly,
            "monthly": self.monthly,
            "posts": self.posts,
            "boosting_susy": self.boosting_susy,
            "susy_budget": self.susy_budget,
            "creators": self.creators,
        }


def parse_workbook(path: Path | str, config: dict[str, Any] | None = None) -> ParseResult:
    cfg = config or load_config()
    accounts = block_accounts(cfg)
    wb = load_workbook(path, data_only=True, read_only=True)
    names = set(wb.sheetnames)
    blocks: list[BlockReport] = []
    issues: list[Issue] = []

    def missing(name: str, what: str) -> None:
        issues.append(
            Issue(source=name, severity="warn", where="workbook", message=f"{what} sheet '{name}' not found")
        )

    weekly_parts, monthly_parts, post_parts = [], [], []
    for spec in cfg["sheets"].get("weekly", []):
        if spec["name"] not in names:
            missing(spec["name"], "weekly")
            continue
        df, rep, iss = parse_report_sheet(
            sheet_rows(wb[spec["name"]]), spec["name"], spec["year"], accounts, None
        )
        weekly_parts.append(df[df["period_type"].eq("cw")] if not df.empty else df)
        if not df.empty and df["period_type"].eq("month").any():
            monthly_parts.append(df[df["period_type"].eq("month")])
        blocks += rep
        issues += iss
    for spec in cfg["sheets"].get("monthly", []):
        if spec["name"] not in names:
            missing(spec["name"], "monthly")
            continue
        df, rep, iss = parse_report_sheet(
            sheet_rows(wb[spec["name"]]), spec["name"], spec["year"], accounts, "month"
        )
        monthly_parts.append(df)
        blocks += rep
        issues += iss
    for spec in cfg["sheets"].get("post_by_post", []):
        if spec["name"] not in names:
            missing(spec["name"], "post-by-post")
            continue
        df, rep, iss = parse_post_by_post(sheet_rows(wb[spec["name"]]), spec["name"], spec["year"], accounts)
        post_parts.append(df)
        blocks += rep
        issues += iss

    susy_spec = cfg["sheets"].get("boosting_susy", {})
    if susy_spec.get("name") in names:
        susy, budget, rep, iss = parse_boosting_susy(
            sheet_rows(wb[susy_spec["name"]]),
            susy_spec["name"],
            susy_spec.get("year", 2026),
            accounts,
            susy_spec.get("budget_col", "E"),
        )
        blocks += rep
        issues += iss
    else:
        missing(str(susy_spec.get("name")), "boosting")
        susy = pd.DataFrame(columns=ACCOUNT_COLS + ["year", "month", "boost_spend_susy", "cell", "n_cells"])
        budget = pd.DataFrame(columns=ACCOUNT_COLS + ["year", "boost_budget", "budget_raw", "cell"])

    cre_spec = cfg["sheets"].get("creators", {})
    if cre_spec.get("name") in names:
        creators, rep, iss = parse_creators(wb[cre_spec["name"]])
        blocks += rep
        issues += iss
    else:
        missing(str(cre_spec.get("name")), "creators")
        creators = pd.DataFrame(
            columns=[
                "name",
                "country",
                "followers",
                "field",
                "account_url",
                "handle",
                "platform",
                "platform_raw",
                "deal",
                "code",
            ]
        )
    wb.close()

    def cat(parts: list[pd.DataFrame], cols: list[str]) -> pd.DataFrame:
        parts = [p for p in parts if not p.empty]
        return pd.concat(parts, ignore_index=True) if parts else pd.DataFrame(columns=cols)

    long_cols = ACCOUNT_COLS[:4] + [
        "sheet",
        "block_title",
        "account",
        "year",
        "period_type",
        "period",
        "metric",
        "value",
        "value_text",
        "cell",
        "n_cells",
    ]
    for b in blocks:
        for w in b.warnings:
            issues.append(
                Issue(
                    source=b.sheet,
                    severity="info",
                    where=f"{b.title} (rows {b.row_start}-{b.row_end})",
                    message=w,
                )
            )
        if b.status == "no_data":
            issues.append(Issue(source=b.sheet, severity="info", where=b.title, message="block has no data"))
    return ParseResult(
        weekly=cat(weekly_parts, long_cols),
        monthly=cat(monthly_parts, long_cols),
        posts=cat(post_parts, POST_COLS),
        boosting_susy=susy,
        susy_budget=budget,
        creators=creators,
        blocks=blocks,
        issues=issues,
        source=str(path),
    )
